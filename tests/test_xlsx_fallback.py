from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from ayman_os_agent.study_sheet import StudySheetManager


def build_workbook(path, today=None):
    today = today or datetime.now(ZoneInfo("Asia/Riyadh")).date().isoformat()
    wb = Workbook()

    config = wb.active
    config.title = "Config"
    config.append(["Key", "Value"])
    config.append(["STUDENT_NAME", "Ayman"])
    config.append(["PARENT_EMAIL", "parent@example.com"])

    courses = wb.create_sheet("Courses")
    courses.append(["Code", "Name", "Risk", "Priority", "WeeklyTargetMin"])
    courses.append(["CS120", "Networks", "HIGH", 1, 120])
    courses.append(["MATH101", "Calculus", "LOW", 3, 60])
    courses.append(["CS230", "Databases", "MEDIUM", 2, 90])

    daily = wb.create_sheet("DailyLog")
    daily.append(["Date", "TotalMin", "Confidence", "HardestCourse", "NeedHelp"])
    daily.append([today, 90, "High", "CS120", "No"])

    wb.create_sheet("Schedule").append(["Day", "Start", "End", "Activity"])
    wb.create_sheet("Alerts").append(["Date", "Severity", "Type", "Message"])

    wb.save(path)
    return path


def test_xlsx_mode_detected(tmp_path):
    path = build_workbook(tmp_path / "study.xlsx")
    manager = StudySheetManager(workbook_path=str(path))
    assert manager.mode() == "xlsx"
    assert manager.is_available() is True


def test_summary_reads_config(tmp_path):
    path = build_workbook(tmp_path / "study.xlsx")
    manager = StudySheetManager(workbook_path=str(path))
    summary = manager.summary()
    assert "Ayman" in summary
    assert "CS120" in summary


def test_risk_ranking(tmp_path):
    path = build_workbook(tmp_path / "study.xlsx")
    manager = StudySheetManager(workbook_path=str(path))
    risk = manager.risk_summary_summary() if hasattr(manager, "risk_summary_summary") else None
    top = manager.top_risk_course()
    assert top["Code"] == "CS120"
    ranked = [c["Code"] for c in manager.high_risk_courses()]
    assert ranked == ["CS120", "CS230"]
    assert risk is None


def test_minutes_logged_today_from_dailylog(tmp_path):
    path = build_workbook(tmp_path / "study.xlsx")
    manager = StudySheetManager(workbook_path=str(path))
    assert manager.minutes_logged_today() == 90


def test_minutes_logged_today_from_sessionlog():
    from conftest import FakeStore

    store = FakeStore(
        tabs={
            "SessionLog": {
                "headers": ["Date", "Time", "CourseCode", "Minutes", "Note"],
                "rows": [
                    ["1900-01-01", "10:00", "CS120", 999, ""],
                    ["1900-01-01", "11:00", "MATH101", 30, ""],  # not today
                ],
            }
        },
        writable=True,
    )
    manager = StudySheetManager(store=store)
    today = datetime.now(ZoneInfo("Asia/Riyadh")).date().isoformat()
    store.tabs["SessionLog"]["rows"][0] = [today, "10:00", "CS120", 40, ""]
    assert manager.mode() == "sheets"
    assert manager.minutes_logged_today() == 40


def test_sessionlog_beats_dailylog_when_larger():
    from conftest import FakeStore

    store = FakeStore(
        tabs={
            "DailyLog": {
                "headers": ["Date", "TotalMin", "Confidence", "HardestCourse", "NeedHelp"],
                "rows": [["1900-01-01", 25, "Low", "", ""]],
            },
            "SessionLog": {
                "headers": ["Date", "Time", "CourseCode", "Minutes", "Note"],
                "rows": [["1900-01-01", "10:00", "CS120", 10, ""]],
            },
        },
        writable=True,
    )
    manager = StudySheetManager(store=store)
    today = datetime.now(ZoneInfo("Asia/Riyadh")).date().isoformat()
    store.tabs["DailyLog"]["rows"][0] = [today, 25, "Low", "", ""]
    store.tabs["SessionLog"]["rows"][0] = [today, "10:00", "CS120", 40, ""]
    # sessions (40) beat the manual daily total (25)
    assert manager.minutes_logged_today() == 40
