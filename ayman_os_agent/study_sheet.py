"""Study workbook access: Google Sheets (primary) or XLSX (read-only fallback).

Exposes normalized records for the tabs: Config, Courses, DailyLog,
SessionLog, Schedule, Alerts.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from .config import get_data_dir, get_study_sheet_path, get_study_sheet_url
from .sheet_store import SheetStore
from .timeutils import parse_date, today_iso

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


def field(record: dict, *aliases: str, default: Any = "") -> Any:
    """Fetch the first matching key from a record (case/space tolerant)."""
    lookup = {str(k).strip().lower().replace(" ", "").replace("_", ""): k for k in record}
    for alias in aliases:
        key = lookup.get(alias.lower().replace(" ", "").replace("_", ""))
        if key is not None and str(record[key]).strip():
            return record[key]
    return default


class StudySheetManager:
    """Read the study tracking data used by the agent."""

    def __init__(self, store: SheetStore | None = None, workbook_path: str | None = None) -> None:
        self.store = store or SheetStore.from_env()
        self.remote_error: str | None = None
        self.workbook_path = self._resolve_xlsx_path(workbook_path)

    # ------------------------------------------------------------ mode

    def mode(self) -> str:
        if self.store.is_writable():
            return "sheets"
        if self.workbook_path and self.workbook_path.exists() and load_workbook is not None:
            return "xlsx"
        if self.store.sheet_url:
            candidate = get_data_dir() / "study-sheet.xlsx"
            if candidate.exists() and load_workbook is not None:
                return "xlsx"
        return "none"

    def is_available(self) -> bool:
        return self.mode() in {"sheets", "xlsx"}

    # ------------------------------------------------------------ xlsx fallback

    def _resolve_xlsx_path(self, workbook_path: str | None):
        if workbook_path:
            return Path(workbook_path).expanduser()
        explicit = get_study_sheet_path()
        if explicit:
            return Path(explicit).expanduser()
        url = get_study_sheet_url()
        if url:
            cached = get_data_dir() / "study-sheet.xlsx"
            try:
                if not cached.exists() or cached.stat().st_size == 0:
                    self._download_remote_workbook(url, cached)
            except (HTTPError, URLError, ValueError, OSError) as exc:
                self.remote_error = f"XLSX fallback download failed: {exc}"
            return cached
        return None

    @staticmethod
    def _download_remote_workbook(url: str, destination) -> None:
        parsed = urlparse(url)
        parts = [part for part in parsed.path.split("/") if part]
        if "spreadsheets" in parts and "d" in parts:
            sheet_id = parts[parts.index("d") + 1]
            url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
        request = Request(url, headers={"User-Agent": "ayman-os-agent/0.3"})
        with urlopen(request, timeout=30) as response:
            content = response.read()
        if not content.startswith(b"PK"):
            raise ValueError("Download did not return an XLSX workbook.")
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)

    def _load_xlsx(self):
        if not self.workbook_path or load_workbook is None:
            return None
        try:
            if not self.workbook_path.exists():
                return None
            return load_workbook(self.workbook_path, data_only=True)
        except Exception:
            return None

    # ------------------------------------------------------------ unified reads

    @staticmethod
    def _normalize(value: Any) -> str:
        if value is None:
            return ""
        if hasattr(value, "isoformat") and not isinstance(value, str):
            return value.isoformat()
        return str(value).strip()

    def _records(self, tab: str) -> list[dict]:
        """Records from Google Sheets (cached) or the XLSX fallback."""
        if self.store.is_writable():
            return self.store.get_records(tab)
        workbook = self._load_xlsx()
        if workbook is None:
            return []
        if tab not in workbook.sheetnames:
            return []
        rows = list(workbook[tab].iter_rows(values_only=True))
        if not rows:
            return []
        headers = [self._normalize(cell) for cell in rows[0]]
        records = []
        for row in rows[1:]:
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            records.append(
                {headers[i]: self._normalize(row[i] if i < len(row) else "") for i in range(len(headers))}
            )
        return records

    def config_map(self) -> dict:
        records = self._records("Config")
        return {field(r, "key", "config", "name"): field(r, "value", "setting") for r in records if field(r, "key", "config", "name")}

    def courses(self) -> list[dict]:
        return self._records("Courses")

    def schedule(self) -> list[dict]:
        return self._records("Schedule")

    def daily_logs(self) -> list[dict]:
        return self._records("DailyLog")

    def session_logs(self) -> list[dict]:
        return self._records("SessionLog")

    def alerts(self) -> list[dict]:
        return self._records("Alerts")

    def latest_daily(self) -> dict | None:
        logs = self.daily_logs()
        return logs[-1] if logs else None

    # ------------------------------------------------------------ derived metrics

    def minutes_logged_today(self) -> float:
        """Minutes logged today: SessionLog sessions + DailyLog TotalMin (whichever larger)."""
        today = today_iso()
        sessions = 0.0
        for record in self.session_logs():
            if parse_date(field(record, "date", "day")) and parse_date(field(record, "date", "day")).isoformat() == today:
                try:
                    sessions += float(field(record, "minutes", "min", "duration", default=0) or 0)
                except (TypeError, ValueError):
                    continue
        daily_total = 0.0
        for record in self.daily_logs():
            day = parse_date(field(record, "date", "day"))
            if day and day.isoformat() == today:
                try:
                    daily_total = max(daily_total, float(field(record, "totalmin", "total_min", "minutes", default=0) or 0))
                except (TypeError, ValueError):
                    continue
        return max(sessions, daily_total)

    def high_risk_courses(self) -> list[dict]:
        entries = []
        for item in self.courses():
            risk = str(field(item, "risk", default="")).upper()
            if risk in {"HIGH", "MEDIUM"}:
                try:
                    priority = float(field(item, "priority", default=999))
                except (TypeError, ValueError):
                    priority = 999.0
                entries.append((priority, item))
        entries.sort(key=lambda pair: pair[0])
        return [item for _, item in entries]

    def top_risk_course(self) -> dict | None:
        ranked = self.high_risk_courses()
        if ranked:
            return ranked[0]
        courses = self.courses()
        return courses[0] if courses else None

    # ------------------------------------------------------------ summary text

    def summary(self) -> str:
        if not self.is_available():
            if self.remote_error:
                return f"تعذر تحميل ورقة الدراسة. {self.remote_error}"
            if self.store.last_error:
                return f"تعذر الوصول إلى Google Sheets: {self.store.last_error}"
            return (
                "ورقة الدراسة غير متصلة. اضبط GOOGLE_SERVICE_ACCOUNT_JSON و GOOGLE_SHEET_ID "
                "(قراءة + كتابة) أو AYMAN_STUDY_SHEET_URL / AYMAN_STUDY_SHEET (قراءة فقط)."
            )

        config_map = self.config_map()
        lines = [
            "## ملخص ورقة الدراسة",
            f"المصدر: {'Google Sheets (مباشر)' if self.mode() == 'sheets' else self.workbook_path}",
            f"اسم الطالب: {field(config_map, 'student_name', default='-')}",
            "",
            "### جدول اليوم والمواد",
        ]

        schedule_preview = []
        for row in self.schedule()[:5]:
            row_summary = " | ".join(str(v) for v in list(row.values())[:4] if str(v).strip())
            if row_summary:
                schedule_preview.append(f"- {row_summary}")
        lines.extend(schedule_preview or ["- لا توجد مواعيد مسجلة في الجدول."])

        lines.extend(["", "### المواد ذات الأولوية العالية"])
        high_risk = self.high_risk_courses()
        if high_risk:
            for item in high_risk[:5]:
                code = field(item, "code", "key", default="-")
                name = field(item, "name", default="-")
                risk = field(item, "risk", default="?")
                priority = field(item, "priority", default="-")
                target = field(item, "weeklytargetmin", "target", default="-")
                lines.append(f"- {code}: {name} | Risk={risk} | Priority={priority} | Target={target} min")
        else:
            lines.append("- لا توجد مواد مطابقة.")

        lines.extend(["", "### آخر سجل يومي"])
        latest = self.latest_daily()
        if latest:
            lines.append(
                f"- TotalMin={field(latest, 'totalmin', 'total_min', default='-')}, "
                f"Confidence={field(latest, 'confidence', default='-')}, "
                f"HardestCourse={field(latest, 'hardestcourse', 'hardest_course', default='-')}, "
                f"NeedHelp={field(latest, 'needhelp', 'need_help', default='-')}"
            )
        else:
            lines.append("- لا يوجد سجل يومي حديث.")

        lines.extend(["", "### آخر تنبيه"])
        alert_rows = self.alerts()
        if alert_rows:
            latest_alert = alert_rows[-1]
            lines.append(
                f"- {field(latest_alert, 'severity', default='-')} | {field(latest_alert, 'type', default='-')} | "
                f"{field(latest_alert, 'message', default='-')}"
            )
        else:
            lines.append("- لا يوجد تنبيه مسجل.")

        return "\n".join(lines)

